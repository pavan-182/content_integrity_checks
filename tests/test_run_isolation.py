"""Synthetic pipeline checks for run isolation and complete failure accounting."""

import json
import os
import subprocess
import sys
from concurrent.futures import ThreadPoolExecutor
from dataclasses import replace
from threading import Barrier
from pathlib import Path

import pysbd
import pytest
from openpyxl import load_workbook

from content_integrity.entity_extraction import EntityExtractor, extract_typed_entities
from content_integrity.pipeline import PipelineConfig, run_pipeline
from content_integrity.validators.context_validator import TruncatedResponseError
from content_integrity.utils import split_sentences


class EntityClient:
    def __init__(self, model_name, barrier=None, fail=False):
        self.model_name = model_name
        self.barrier = barrier
        self.fail = fail
        self.calls = 0

    def complete(self, **kwargs):
        self.calls += 1
        if self.barrier and self.calls == 1:
            self.barrier.wait(timeout=10)
        if self.fail and "Badmarker" in kwargs["user"]:
            raise RuntimeError("Synthetic entity gateway failure for Badmarker")
        return json.dumps({"entities": [{"text": "Novelmarker", "type": self.model_name}]})


def config(tmp_path):
    source = tmp_path / "input"
    source.mkdir()
    for name, text in (("GOOD", "Novelmarker expression was measured."), ("BAD", "Badmarker expression was measured.")):
        (source / f"{name}.xml").write_text(
            f'<article><front><article-meta><article-id pub-id-type="manuscript">{name}</article-id>'
            f'<title-group><article-title>Synthetic study</article-title></title-group>'
            f'<abstract><p>{text}</p></abstract></article-meta></front></article>', encoding="utf-8",
        )
    dictionary = tmp_path / "dictionary.csv"
    dictionary.write_text("Fingerprint - Tortured Phrase,Expected Text,Nb Retrieved Papers\n", encoding="utf-8")
    return PipelineConfig(source, tmp_path / "output", dictionary)


def summary(item):
    return item["checks"][0]["result"]["supporting_data"][0]


def test_sequential_and_concurrent_pipeline_runs_own_clients_and_counts(tmp_path):
    settings = config(tmp_path)
    barrier = Barrier(2)
    clients = [EntityClient("gene", barrier), EntityClient("protein", barrier)]
    with ThreadPoolExecutor(max_workers=2) as pool:
        futures = [pool.submit(run_pipeline, replace(settings, output_dir=tmp_path / str(i)), llm_client=client)
                   for i, client in enumerate(clients)]
        results = [future.result(timeout=30) for future in futures]
    for client, result in zip(clients, results):
        metadata = dict(result.run_metadata_rows)
        assert metadata["entity_model_id"] == client.model_name
        assert metadata["entity_model_inference_count"] == client.calls == 2
        assert not result.operational_issues
    offline = run_pipeline(replace(settings, offline=True))
    assert dict(offline.run_metadata_rows)["entity_model_inference_count"] == 0
    assert not offline.operational_issues
    another = EntityClient("disease")
    run_pipeline(replace(settings, output_dir=tmp_path / "another"), llm_client=another)
    assert another.calls == 2
    assert not extract_typed_entities("Novelmarker expression")


def test_entity_spans_are_isolated_and_reused_within_each_run():
    left, right = EntityExtractor(EntityClient("gene")), EntityExtractor(EntityClient("protein"))
    for extractor, expected in ((left, "gene"), (right, "protein"), (left, "gene")):
        entities = extract_typed_entities("Novelmarker", extractor=extractor)
        assert [(entity.text, entity.entity_type) for entity in entities] == [("Novelmarker", expected)]
        assert extractor.inference_count == 1


def test_one_failure_preserves_other_records_and_reconciles_json_workbook(tmp_path):
    settings = config(tmp_path)
    (settings.input_dir / "broken.xml").write_text("<", encoding="utf-8")
    result = run_pipeline(settings, llm_client=EntityClient("gene", fail=True))
    report = json.loads(result.output_paths["content_integrity_json"].read_text())
    assert len(result.records) == len(report) == 3
    by_id = {item["abstract_id"]: summary(item) for item in report.values()}
    assert by_id["GOOD"]["processing_status"] == "successful"
    assert by_id["BAD"]["processing_status"] == "failed"
    assert "Badmarker" in by_id["BAD"]["operational_issues"][0]["message"]
    assert sum(item["processing_status"] == "failed" for item in by_id.values()) == 2
    template = next(check for check in report["BAD"]["checks"] if check["check_name"] == "template_detection")
    assert template["result"]["level"] == "UNKNOWN"
    with result.output_paths["workbook"].open("rb") as handle:
        workbook = load_workbook(handle, read_only=True, data_only=True)
        rows = workbook["All Abstracts"].iter_rows(values_only=True)
        headers = next(rows)
        rows = [dict(zip(headers, row)) for row in rows]
        assert len(rows) == len(report)
        for row in rows:
            assert row["Processing Status"] == by_id[row["Abstract ID"]]["processing_status"]
        workbook.close()


def test_truncated_unbroken_chunk_splits_and_counts_attempts():
    class TruncatingClient:
        def complete(self, **kwargs):
            if len(kwargs["user"]) > 120:
                raise TruncatedResponseError("synthetic token limit")
            return '{"entities": []}'

    extractor = EntityExtractor(TruncatingClient())
    assert extractor.spans("x" * 240) == ()
    assert extractor.inference_count == 3


def test_concurrent_sentence_splitting_preserves_each_input(monkeypatch):
    barrier = Barrier(2)
    original = pysbd.Segmenter.processor

    def interleave(segmenter, text):
        barrier.wait(timeout=10)
        return original(segmenter, text)

    monkeypatch.setattr(pysbd.Segmenter, "processor", interleave)
    inputs = ["First synthetic sentence.", "Different synthetic sentence."]
    with ThreadPoolExecutor(max_workers=2) as pool:
        assert list(pool.map(split_sentences, inputs)) == [[text] for text in inputs]


@pytest.mark.parametrize("kind", ["missing", "empty", "file"])
def test_cli_rejects_invalid_input_before_writing_outputs(tmp_path, kind):
    source = tmp_path / "input"
    if kind == "empty":
        source.mkdir()
    elif kind == "file":
        source.write_text("<article/>", encoding="utf-8")
    output = tmp_path / "output"
    root = Path(__file__).resolve().parents[1]
    result = subprocess.run(
        [sys.executable, str(root / "scripts/run_pipeline.py"), "--offline", "--input-dir", str(source), "--output-dir", str(output)],
        capture_output=True, text=True, timeout=20,
    )
    assert result.returncode != 0
    assert "Input directory" in result.stderr
    assert "ZeroDivisionError" not in result.stderr
    assert not output.exists()


def test_cli_reports_are_identical_across_hash_seeds(tmp_path):
    settings = config(tmp_path)
    settings.tortured_dictionary_path.write_text(
        'Fingerprint - Tortured Phrase,Expected Text,Nb Retrieved Papers\n'
        '"nervous network","neural network",1\n'
        '"counterfeit consciousness","artificial intelligence",1\n'
        '"neural organization","neural network",1\n', encoding="utf-8",
    )
    (settings.input_dir / "GOOD.xml").write_text(
        '<article><front><article-meta><article-id pub-id-type="manuscript">GOOD</article-id>'
        '<abstract><p>A nervous network uses counterfeit consciousness and a neural organization.</p></abstract>'
        '</article-meta></front></article>', encoding="utf-8",
    )
    root = Path(__file__).resolve().parents[1]
    reports = []
    for seed in ("1", "7", "42"):
        output = tmp_path / seed
        subprocess.run(
            [sys.executable, str(root / "scripts/run_pipeline.py"), "--offline", "--input-dir", str(settings.input_dir),
             "--tortured-dictionary", str(settings.tortured_dictionary_path), "--output-dir", str(output)],
            env={**os.environ, "PYTHONHASHSEED": seed}, capture_output=True, text=True, check=True, timeout=30,
        )
        reports.append(json.loads((output / "content_integrity_results.json").read_text()))
    assert reports[0] == reports[1] == reports[2]
