"""Phase 2 reliability: record isolation, reconciliation, atomic promotion, checkpoint and resume.

Inputs are small synthetic datasets from scripts/generate_load_dataset.py; nothing here measures
detector accuracy.
"""

from __future__ import annotations

import hashlib
import json
import os
import signal
import subprocess
import sys
import time
from dataclasses import replace
from pathlib import Path

import pytest
from openpyxl import load_workbook

import content_integrity.detectors.numerical_contradiction as numerical
import content_integrity.pipeline as pipeline
from content_integrity.checkpoint import IncompatibleCheckpointError
from content_integrity.pipeline import PipelineConfig, run_pipeline
from content_integrity.reconciliation import ReconciliationError, reconcile_outputs
from content_integrity.validators.context_validator import IntelliHubGPTOSSClient
from scripts.fake_gpt_oss_gateway import API_KEY, MODEL, FakeGateway, GatewayBehaviour
from scripts.generate_load_dataset import SyntheticAbstract, _bundle_xml, generate

ROOT = Path(__file__).resolve().parents[1]


@pytest.fixture
def dictionary(tmp_path):
    path = tmp_path / "dictionary.csv"
    path.write_text(
        'Fingerprint - Tortured Phrase,Expected Text,Nb Retrieved Papers\n'
        '"counterfeit consciousness","artificial intelligence",1\n'
        '"irregular woodland","random forest",1\n',
        encoding="utf-8",
    )
    return path


def _config(tmp_path: Path, dictionary: Path, profile: str = "ci", records: int | None = 16, **overrides) -> PipelineConfig:
    source = tmp_path / f"input-{profile}-{records}"
    if not source.exists():
        generate(profile, source, records=records)
    return PipelineConfig(source, tmp_path / "output", dictionary, offline=True, **overrides)


def _sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _checkpoint_stages(output: Path) -> list[str]:
    return json.loads((output / ".checkpoint" / "checkpoint.json").read_text())["completed_stages"]


def _summaries(path: Path) -> dict[str, dict]:
    report = json.loads(path.read_text())
    return {entry["abstract_id"]: entry["checks"][0]["result"]["supporting_data"][0] for entry in report.values()}


def test_every_malformed_input_gets_a_terminal_status_and_reports_reconcile(tmp_path, dictionary):
    settings = _config(tmp_path, dictionary, profile="malformed", records=None)
    manifest = json.loads((settings.input_dir / "manifest.json").read_text())
    result = run_pipeline(settings)
    records = result.run_summary["records"]
    assert result.run_summary["reconciled"], result.run_summary["checks"]
    assert records["total_input"] == manifest["expected"]["input_records"]
    assert records["completed"] + records["completed_with_findings"] + records["failed"] + records["skipped"] == records["total_input"]
    assert records["skipped"] == manifest["expected"]["duplicate_records"]
    categories = result.run_summary["failures_by_stage_and_category"]
    assert categories["record_identity:duplicate_identifier"] == 2  # same record ID and DOI in two files
    assert {failure["error_category"] for failure in result.run_summary["failures"]} <= {"invalid_input", "duplicate_identifier"}
    assert all(failure["retryable"] is False for failure in result.run_summary["failures"])
    summaries = _summaries(result.output_paths["content_integrity_json"])
    assert {summary["record_status"] for summary in summaries.values()} <= {"completed", "completed_with_findings", "failed"}


def test_one_record_raising_in_a_whole_batch_detector_does_not_remove_the_check_for_others(tmp_path, dictionary, monkeypatch):
    settings = _config(tmp_path, dictionary)
    original = numerical.extract_numerical_claims
    victim = "synci-000003"

    def failing(record):
        if record.record_id == victim:
            raise ValueError("synthetic extraction failure")
        return original(record)

    monkeypatch.setattr(numerical, "extract_numerical_claims", failing)
    result = run_pipeline(settings)
    failures = result.run_summary["failures"]
    assert result.run_summary["reconciled"]
    assert [(item["record_id"], item["stage"], item["error_category"]) for item in failures] == [
        (victim, "numerical_contradiction", "processing_error"),
    ]
    assert result.run_summary["records"]["failed"] == 1


def test_interrupted_template_stage_resumes_without_recomputing_and_matches_a_clean_run(tmp_path, dictionary, monkeypatch):
    clean = run_pipeline(replace(_config(tmp_path, dictionary), output_dir=tmp_path / "clean"))
    settings = _config(tmp_path, dictionary)

    def interrupt(*args, **kwargs):
        raise KeyboardInterrupt

    with monkeypatch.context() as patched:
        patched.setattr(pipeline, "detect_entity_normalized_templates", interrupt)
        with pytest.raises(KeyboardInterrupt):
            run_pipeline(settings)
    assert _checkpoint_stages(settings.output_dir) == ["record_checks", "template_features"]
    assert not (settings.output_dir / "content_integrity_results.json").exists()

    monkeypatch.setattr(pipeline, "build_template_features", lambda *a, **k: pytest.fail("features were recomputed"))
    monkeypatch.setattr(pipeline, "detect_tortured_phrases", lambda *a, **k: pytest.fail("record checks were recomputed"))
    resumed = run_pipeline(settings)
    assert _sha(resumed.output_paths["content_integrity_json"]) == _sha(clean.output_paths["content_integrity_json"])
    assert resumed.run_metrics["resumed_from_run_ids"]
    assert [stage["stage"] for stage in resumed.run_metrics["stages"] if stage.get("restored_from_checkpoint")] == [
        "record_checks", "template_features",
    ]
    assert not (settings.output_dir / ".checkpoint").exists()
    assert resumed.run_summary["records"] == clean.run_summary["records"]


def test_output_write_failure_keeps_previous_reports_and_the_retry_resumes_at_reporting(tmp_path, dictionary, monkeypatch):
    settings = _config(tmp_path, dictionary)
    first = run_pipeline(settings)
    previous = {name: _sha(path) for name, path in first.output_paths.items()}

    def disk_full(*args, **kwargs):
        raise OSError(28, "No space left on device")

    with monkeypatch.context() as patched:
        patched.setattr(pipeline, "write_workbook", disk_full)
        with pytest.raises(OSError):
            run_pipeline(settings)
    assert {name: _sha(path) for name, path in first.output_paths.items()} == previous
    assert not list(settings.output_dir.glob(".staging-*"))
    assert _checkpoint_stages(settings.output_dir) == list(pipeline.STAGES)

    monkeypatch.setattr(pipeline, "detect_exact_text_reuse", lambda *a, **k: pytest.fail("template pairs were recomputed"))
    retried = run_pipeline(settings)
    assert _sha(retried.output_paths["content_integrity_json"]) == previous["content_integrity_json"]
    assert retried.run_summary["run_id"] != first.run_summary["run_id"]


def test_reconciliation_failure_never_replaces_previous_reports(tmp_path, dictionary, monkeypatch):
    settings = _config(tmp_path, dictionary)
    first = run_pipeline(settings)
    previous = {name: _sha(path) for name, path in first.output_paths.items()}
    original = pipeline.reconcile_outputs

    def broken(**kwargs):
        result = original(**kwargs)
        return {**result, "reconciled": False, "checks": [{"name": "injected", "passed": False, "actual": 1, "expected": 0}]}

    monkeypatch.setattr(pipeline, "reconcile_outputs", broken)
    with pytest.raises(ReconciliationError, match="injected"):
        run_pipeline(settings)
    assert {name: _sha(path) for name, path in first.output_paths.items()} == previous
    assert len(list(settings.output_dir.glob(".staging-*"))) == 1  # retained for inspection


def test_reconciliation_detects_a_workbook_that_lost_a_record(tmp_path, dictionary):
    result = run_pipeline(_config(tmp_path, dictionary))
    workbook_path = result.output_paths["workbook"]
    workbook = load_workbook(workbook_path)
    workbook["All Abstracts"].delete_rows(workbook["All Abstracts"].max_row)
    workbook.save(workbook_path)
    checked = reconcile_outputs(
        record_ids=[record.record_id for record in result.records],
        input_record_count=len(result.records),
        skipped=[],
        reportable_findings=[finding for finding in result.findings if finding.detector_type != "nonsense_candidate"],
        reviewable_pair_count=result.run_summary["template"]["reviewable_pairs"],
        json_path=result.output_paths["content_integrity_json"],
        workbook_path=workbook_path,
    )
    failed = {check["name"] for check in checked["checks"] if not check["passed"]}
    assert not checked["reconciled"]
    assert {"workbook_abstracts_equal_retained_records", "workbook_ids_match_retained_records"} <= failed


def test_incompatible_or_corrupt_checkpoints_are_rejected_until_discarded(tmp_path, dictionary, monkeypatch):
    settings = _config(tmp_path, dictionary)
    with monkeypatch.context() as patched:
        patched.setattr(pipeline, "detect_exact_text_reuse", lambda *a, **k: (_ for _ in ()).throw(KeyboardInterrupt))
        with pytest.raises(KeyboardInterrupt):
            run_pipeline(settings)
    dictionary.write_text(dictionary.read_text() + '"bosom peril","breast cancer",1\n', encoding="utf-8")
    with pytest.raises(IncompatibleCheckpointError, match="tortured_dictionary_sha256"):
        run_pipeline(settings)

    index = json.loads((settings.output_dir / ".checkpoint" / "checkpoint.json").read_text())
    index["fingerprint"]["tortured_dictionary_sha256"] = hashlib.sha256(dictionary.read_bytes()).hexdigest()
    (settings.output_dir / ".checkpoint" / "checkpoint.json").write_text(json.dumps(index))
    (settings.output_dir / ".checkpoint" / index["state_file"]).write_bytes(b"corrupted")
    with pytest.raises(IncompatibleCheckpointError, match="corrupt"):
        run_pipeline(settings)

    result = run_pipeline(replace(settings, discard_checkpoint=True))
    assert result.run_summary["reconciled"] and not result.run_metrics["resumed_from_run_ids"]


def _gateway_client(gateway: FakeGateway, **overrides) -> IntelliHubGPTOSSClient:
    settings = dict(api_key=API_KEY, base_url=gateway.base_url, model_name=MODEL, backoff_seconds=0.01, max_backoff_seconds=0.02, timeout_seconds=2.0)
    return IntelliHubGPTOSSClient(**{**settings, **overrides})


def test_transient_gateway_failures_are_retried_and_every_record_completes(tmp_path, dictionary):
    behaviour = GatewayBehaviour(failure_rate=0.25, failure_modes=("429", "500", "503", "reset", "malformed"))
    with FakeGateway(behaviour) as gateway:
        client = _gateway_client(gateway, max_attempts=6)
        result = run_pipeline(replace(_config(tmp_path, dictionary), offline=False), llm_client=client)
    assert result.run_summary["reconciled"]
    assert result.run_summary["records"]["failed"] == 0
    assert client.call_stats.retry_count > 0 and client.call_stats.failure_count == 0
    assert result.run_metrics["gateway"]["max_in_flight_requests"] <= client.max_concurrent_requests


def test_gateway_outage_fails_records_as_retryable_without_unbounded_requests(tmp_path, dictionary):
    with FakeGateway(GatewayBehaviour(force_status=503)) as gateway:
        client = _gateway_client(gateway, max_attempts=2, circuit_failure_threshold=3, circuit_cooldown_seconds=60)
        result = run_pipeline(replace(_config(tmp_path, dictionary), offline=False), llm_client=client)
        requests = gateway.counters.snapshot()["requests"]
    summary = result.run_summary
    assert summary["reconciled"]
    assert summary["records"]["failed"] == summary["records"]["total_input"]
    categories = {failure["error_category"] for failure in summary["failures"]}
    assert categories == {"server_error", "circuit_open"}
    assert all(failure["retryable"] for failure in summary["failures"])
    # The circuit opened after three exhausted requests; later records failed fast without traffic.
    assert requests <= 3 * 2 + client.max_concurrent_requests * 2


def test_killed_process_resumes_without_repeating_completed_model_calls(tmp_path, dictionary):
    settings = _config(tmp_path, dictionary, records=12)
    command = [sys.executable, str(ROOT / "scripts/run_pipeline.py"), "--input-dir", str(settings.input_dir),
               "--tortured-dictionary", str(dictionary), "--llm-max-concurrency", "2", "--quiet"]
    with FakeGateway(GatewayBehaviour(latency_seconds=0.05)) as gateway:
        environment = {**os.environ, **gateway.env()}
        reference = subprocess.run([*command, "--output-dir", str(tmp_path / "reference")], env=environment, capture_output=True, text=True, timeout=120)
        assert reference.returncode == 0, reference.stderr
        clean_requests = gateway.counters.snapshot()["requests"]

        output = tmp_path / "interrupted"
        process = subprocess.Popen([*command, "--output-dir", str(output)], env=environment, stdout=subprocess.DEVNULL, stderr=subprocess.PIPE)
        deadline = time.monotonic() + 60
        while gateway.counters.snapshot()["requests"] < clean_requests * 1.5 and time.monotonic() < deadline:
            time.sleep(0.02)
        process.send_signal(signal.SIGKILL)
        process.wait(timeout=30)
        before_resume = gateway.counters.snapshot()["requests"]
        assert _checkpoint_stages(output) == ["record_checks"]
        assert before_resume - clean_requests < clean_requests  # killed mid-extraction

        resumed = subprocess.run([*command, "--output-dir", str(output)], env=environment, capture_output=True, text=True, timeout=120)
        assert resumed.returncode == 0, resumed.stderr
        resumed_requests = gateway.counters.snapshot()["requests"] - before_resume
    completed_before_kill = before_resume - clean_requests
    assert resumed_requests <= clean_requests - completed_before_kill + 2  # at most the in-flight calls repeat
    assert _sha(output / "content_integrity_results.json") == _sha(tmp_path / "reference" / "content_integrity_results.json")
    summary = json.loads((output / "run_summary.json").read_text())
    assert summary["reconciled"] and summary["records"]["failed"] == 0


def test_uncached_trial_ids_are_not_failures_when_registry_lookup_is_not_requested(tmp_path, dictionary):
    settings = _config(tmp_path, dictionary)
    trial = SyntheticAbstract(
        record_id="syntrial-000001", title="Synthetic registered trial abstract", subject="Care Delivery/Models of Care",
        sections=(("Methods", "This study was registered as NCT01234567 and enrolled synthetic patients."),
                  ("Results", "Synthetic outcomes improved across synthetic sites.")),
    )
    (settings.input_dir / "trial.xml").write_text(_bundle_xml([trial]), encoding="utf-8")
    result = run_pipeline(settings)
    assert [record.record_id for record in result.records if record.trial_ids] == ["syntrial-000001"]
    assert result.run_summary["records"]["failed"] == 0
    assert not any(issue.component == "clinical_trial_registry" for issue in result.operational_issues)
